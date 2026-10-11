"""Exportación a Excel.

Cada listado tenía su propia versión del mismo bloque: crear el libro, escribir la
fila de encabezados en negrita, volcar las filas y devolver un StreamingResponse.
Con siete listados más por exportar, esa duplicación se vuelve un problema de
mantenimiento: cualquier cambio (por ejemplo el ancho de columna) hay que hacerlo
en once lugares.

Acá queda una sola función. Un endpoint de exportación pasa a ser:

    return hoja_excel("Máquinas", ["Cliente", "Planta"], filas, "maquinas.xlsx")
"""

import datetime
import io

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _texto(valor):
    """Convierte un valor a algo que Excel muestre bien.

    Los None van vacíos y no como la cadena "None", que es lo que pasaba al
    concatenar sin cuidado.
    """
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, (datetime.date, datetime.datetime)):
        return valor.strftime("%d/%m/%Y")
    return str(valor)


def hoja_excel(titulo, encabezados, filas, nombre_archivo):
    """Planilla de una hoja con encabezados en negrita y columnas ajustadas."""
    wb = Workbook()
    ws = wb.active
    # Excel corta los títulos de hoja en 31 caracteres y no acepta algunos símbolos.
    ws.title = str(titulo)[:31]

    ws.append([str(e) for e in encabezados])
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.alignment = Alignment(vertical="center")

    for fila in filas:
        ws.append([_texto(v) for v in fila])

    # Ancho por columna, acotado para que una observación larga no deje el resto
    # de las columnas fuera de pantalla.
    for i, encabezado in enumerate(encabezados, start=1):
        letra = get_column_letter(i)
        mayor = len(str(encabezado))
        for fila in filas:
            if i - 1 < len(fila):
                mayor = max(mayor, len(_texto(fila[i - 1])))
        ws.column_dimensions[letra].width = min(mayor + 3, 60)

    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    if not nombre_archivo.endswith(".xlsx"):
        nombre_archivo += ".xlsx"

    return StreamingResponse(
        buffer,
        media_type=TIPO_XLSX,
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"},
    )


def nombre_con_fecha(base):
    """`maquinas-2026-10-11.xlsx`, para no pisar exportaciones anteriores."""
    return f"{base}-{datetime.date.today().isoformat()}.xlsx"
