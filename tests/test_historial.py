"""Pruebas de /historial-at: paginación, filtros y exportación.

La tabla volcaba todas las asignaciones del rango sin límite, y no tenía filtro
por estado ni exportación.
"""

import io

from openpyxl import load_workbook

import models


def _seed(db, cantidad=8):
    """Períodos de asignación: la mitad vigentes, la mitad cerrados."""
    import datetime

    db.query(models.ATAssignment).delete()
    db.commit()
    for i in range(cantidad):
        vigente = i % 2 == 0
        db.add(models.ATAssignment(
            customer_name=f"Cliente {i:02d}",
            technical_advisor="ING. GOMEZ" if i < cantidad // 2 else "ING. TORRES",
            standard_visits=i,
            premium_visits=0,
            status="Activo",
            valid_from=datetime.datetime(2026, 1, 1) + datetime.timedelta(days=i),
            valid_to=None if vigente else datetime.datetime(2026, 6, 1),
        ))
    db.commit()


def _filas(html):
    return html.count('<tr class="hover:bg')


def test_pagina_inicial_muestra_solo_el_tamano_de_pagina(client, db_session, admin_cookie):
    _seed(db_session, 30)
    html = client.get("/historial-at").text
    assert _filas(html) == 25, "la página debe traer 25 filas, no todas"
    assert "30 períodos" in html


def test_paginacion_conserva_los_filtros(client, db_session, admin_cookie):
    _seed(db_session, 30)
    html = client.get("/historial-at?cliente=Cliente&page=2").text
    assert "Página 2 de 2" in html
    # los enlaces de paginación deben arrastrar el filtro
    assert "cliente=Cliente" in html
    assert _filas(html) == 5, "la segunda página debería tener las 5 restantes"


def test_filtro_por_vigencia(client, db_session, admin_cookie):
    _seed(db_session, 8)
    vigentes = client.get("/historial-at?estado=vigentes").text
    cerrados = client.get("/historial-at?estado=cerrados").text
    assert "4 períodos" in vigentes
    assert "4 períodos" in cerrados


def test_filtro_por_at(client, db_session, admin_cookie):
    _seed(db_session, 8)
    html = client.get("/historial-at?at=GOMEZ").text
    assert "4 períodos" in html
    # Sólo se miran las filas de la tabla: el nombre del otro AT puede aparecer
    # legítimamente en el panel de distribución "Por AT".
    assert _filas(html) == 4, "el filtro por AT no recortó las filas"


def test_filtro_por_cliente(client, db_session, admin_cookie):
    _seed(db_session, 8)
    html = client.get("/historial-at?cliente=Cliente 03").text
    assert "1 período" in html


def test_exportar_historial_a_excel(client, db_session, admin_cookie):
    _seed(db_session, 5)
    r = client.get("/historial-at/exportar-excel")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert r.content[:2] == b"PK"

    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.cell(row=1, column=1).value == "Cliente"
    assert ws.max_row == 6, "5 filas de datos más el encabezado"
    # la columna "Válido hasta" debe decir "Vigente" cuando no hay cierre
    valores = [ws.cell(row=f, column=5).value for f in range(2, 7)]
    assert "Vigente" in valores


def test_exportar_respeta_el_filtro(client, db_session, admin_cookie):
    _seed(db_session, 8)
    r = client.get("/historial-at/exportar-excel?at=GOMEZ")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == 5, "4 filas de ese AT más el encabezado"
    asesores = {ws.cell(row=f, column=2).value for f in range(2, ws.max_row + 1)}
    assert asesores == {"ING. GOMEZ"}


def test_sin_datos_muestra_mensaje_adecuado(client, db_session, admin_cookie):
    db_session.query(models.ATAssignment).delete()
    db_session.commit()
    html = client.get("/historial-at").text
    assert "No hay períodos registrados" in html

    # con un filtro activo, el mensaje debe hablar de filtros
    html = client.get("/historial-at?cliente=Nadie").text
    assert "coincide con los filtros" in html
