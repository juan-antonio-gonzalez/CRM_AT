"""Pruebas de la importación masiva desde Excel.

La importación escribe datos reales, así que lo que más importa verificar no es
que importe bien, sino que **no importe lo que no debe**: filas sin el campo
obligatorio, clientes repetidos, máquinas de un cliente inexistente, ni la misma
planilla dos veces.
"""

import io
import re

import models
import pytest
from openpyxl import Workbook

from services import importacion


def _csrf(client, ruta="/importar"):
    """Token CSRF del formulario. Sin él, el middleware rechaza el POST."""
    html = client.get(ruta).text
    m = re.search(r'name="csrf_token" value="([^"]*)"', html)
    return m.group(1) if m else ""


def _planilla(encabezados, filas):
    """Arma un .xlsx en memoria con los encabezados y las filas dadas."""
    wb = Workbook()
    ws = wb.active
    ws.append(list(encabezados))
    for f in filas:
        ws.append(list(f))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


ENCABEZADOS_CLIENTES = ["Nombre", "Planta", "Sector industrial", "Equipo",
                        "Asesor comercial", "Asistente técnico", "Diseñador",
                        "Dirección", "Ciudad", "Estado", "Observaciones"]

ENCABEZADOS_MAQUINAS = ["Cliente", "Planta", "Marca", "Modelo", "Tipo",
                        "Nº de serie", "Línea", "Estado"]


# ── Lectura de la planilla ─────────────────────────────────────────────────
def test_lee_una_planilla_de_clientes():
    contenido = _planilla(ENCABEZADOS_CLIENTES, [
        ["Cartocor SA", "Planta 1", "Corrugado", "", "Carlos", "ING. GOMEZ",
         "Diseño", "Calle 1", "Rosario", "Activo", ""],
        ["Envalora SRL", "Planta 2", "Flexible", "", "Ana", "ING. DIAZ",
         "Diseño", "Calle 2", "Rosario", "Activo", ""],
    ])
    filas, error = importacion.leer_planilla(contenido, "clientes")
    assert error is None
    assert len(filas) == 2
    assert filas[0]["name"] == "Cartocor SA"
    # La fila de encabezados es la 1, así que los datos empiezan en la 2.
    assert filas[0]["_fila"] == 2
    assert filas[1]["_fila"] == 3


def test_los_encabezados_se_buscan_por_nombre_no_por_posicion():
    """Reordenar las columnas no rompe la importación."""
    contenido = _planilla(["Ciudad", "NOMBRE", "Planta"], [["Rosario", "Cartocor SA", "P1"]])
    filas, error = importacion.leer_planilla(contenido, "clientes")
    assert error is None
    assert filas[0]["name"] == "Cartocor SA"
    assert filas[0]["city"] == "Rosario"


def test_los_acentos_y_mayusculas_no_importan():
    """«Nº de serie» y «N DE SERIE» tienen que encontrar la misma columna."""
    contenido = _planilla(["Cliente", "N DE SERIE"], [["Cartocor", "SN-1"]])
    filas, error = importacion.leer_planilla(contenido, "maquinas")
    assert error is None
    assert filas[0]["serial_number"] == "SN-1"


def test_falta_una_columna_obligatoria():
    contenido = _planilla(["Planta", "Ciudad"], [["P1", "Rosario"]])
    filas, error = importacion.leer_planilla(contenido, "clientes")
    assert filas == []
    assert "Nombre" in error


def test_archivo_ilegible():
    filas, error = importacion.leer_planilla(b"esto no es un xlsx", "clientes")
    assert filas == []
    assert "Excel" in error or "leer" in error


def test_planilla_sin_filas_de_datos():
    contenido = _planilla(ENCABEZADOS_CLIENTES, [])
    filas, error = importacion.leer_planilla(contenido, "clientes")
    assert filas == []
    assert error and "filas de datos" in error


def test_las_filas_vacias_se_ignoran():
    contenido = _planilla(ENCABEZADOS_CLIENTES, [
        ["Cartocor SA", "", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", ""],
        [None, None, None, None, None, None, None, None, None, None, None],
    ])
    filas, error = importacion.leer_planilla(contenido, "clientes")
    assert error is None
    assert len(filas) == 1


# ── Validación ─────────────────────────────────────────────────────────────
def test_cliente_sin_nombre_se_rechaza(db_session):
    filas = [{"_fila": 2, "name": "", "plant": "P1"},
             {"_fila": 3, "name": "Válido", "plant": "P2"}]
    listas, errores, _avisos = importacion.validar(db_session, "clientes", filas)
    assert len(listas) == 1
    assert listas[0]["name"] == "Válido"
    assert len(errores) == 1
    assert "nombre" in errores[0]["motivo"].lower()


def test_cliente_repetido_en_la_base(db_session):
    db_session.add(models.Customer(name="Cartocor SA", is_active=True))
    db_session.commit()
    filas = [{"_fila": 2, "name": "cartocor sa", "plant": ""}]
    listas, errores, _avisos = importacion.validar(db_session, "clientes", filas)
    # La comparación ignora mayúsculas: es el mismo cliente.
    assert listas == []
    assert "Ya existe" in errores[0]["motivo"]


def test_cliente_repetido_en_la_misma_planilla(db_session):
    filas = [{"_fila": 2, "name": "Nuevo SA", "plant": ""},
             {"_fila": 3, "name": "Nuevo SA", "plant": ""}]
    listas, errores, _avisos = importacion.validar(db_session, "clientes", filas)
    assert len(listas) == 1
    assert len(errores) == 1
    assert "repetido" in errores[0]["motivo"]


def test_maquina_de_cliente_inexistente(db_session):
    filas = [{"_fila": 2, "_cliente": "No Existe SRL", "serial_number": "SN-1"}]
    listas, errores, _avisos = importacion.validar(db_session, "maquinas", filas)
    assert listas == []
    assert "No existe un cliente" in errores[0]["motivo"]


def test_maquina_de_cliente_existente_resuelve_el_id(db_session):
    db_session.add(models.Customer(name="Cartocor SA", is_active=True))
    db_session.commit()
    filas = [{"_fila": 2, "_cliente": "cartocor sa", "serial_number": "SN-1",
              "brand": "Bobst"}]
    listas, errores, _avisos = importacion.validar(db_session, "maquinas", filas)
    assert not errores
    assert listas[0]["customer_id"] == db_session.query(models.Customer.id).first()[0]


def test_numero_de_serie_repetido(db_session):
    c = models.Customer(name="Cartocor SA", is_active=True)
    db_session.add(c)
    db_session.commit()
    db_session.add(models.Machine(customer_id=c.id, serial_number="SN-1", is_active=True))
    db_session.commit()

    filas = [{"_fila": 2, "_cliente": "Cartocor SA", "serial_number": "SN-1"}]
    listas, errores, _avisos = importacion.validar(db_session, "maquinas", filas)
    assert listas == []
    assert "Ya existe una máquina" in errores[0]["motivo"]


# ── Importación ────────────────────────────────────────────────────────────
def test_importar_clientes_crea_los_registros(db_session):
    filas = [{"_fila": 2, "name": "Uno SA", "plant": "P1", "city": "Rosario"},
             {"_fila": 3, "name": "Dos SA", "plant": "P2", "city": "Córdoba"},
             {"_fila": 4, "name": "", "plant": "P3"}]
    creados, omitidos, errores = importacion.importar(db_session, "clientes", filas)
    assert creados == 2
    assert omitidos == 1
    assert db_session.query(models.Customer).count() == 2
    # Los que entraron quedan activos y con el relevamiento pendiente.
    c = db_session.query(models.Customer).filter_by(name="Uno SA").first()
    assert c.is_active is True
    assert c.relevamiento_status == "Pendiente"


def test_importar_dos_veces_la_misma_planilla_no_duplica(db_session):
    """La segunda pasada no debe crear nada: ya existen."""
    filas = [{"_fila": 2, "name": "Uno SA", "plant": "P1"}]

    creados1, _, _ = importacion.importar(db_session, "clientes", filas)
    assert creados1 == 1

    creados2, omitidos2, errores2 = importacion.importar(db_session, "clientes", filas)
    assert creados2 == 0
    assert omitidos2 == 1
    assert db_session.query(models.Customer).count() == 1
    assert "Ya existe" in errores2[0]["motivo"]


def test_importar_maquinas_con_su_cliente(db_session):
    filas_c = [{"_fila": 2, "name": "Cartocor SA", "plant": "P1"}]
    importacion.importar(db_session, "clientes", filas_c)

    filas_m = [{"_fila": 2, "_cliente": "Cartocor SA", "serial_number": "SN-1",
                "brand": "Bobst", "machine_type": "Impresora"}]
    creados, _, errores = importacion.importar(db_session, "maquinas", filas_m)
    assert creados == 1 and not errores

    m = db_session.query(models.Machine).first()
    assert m.brand == "Bobst"
    assert m.serial_number == "SN-1"


def test_la_plantilla_tiene_los_encabezados_correctos():
    """La plantilla descargable sirve para completar y volver a subir."""
    from openpyxl import load_workbook

    respuesta = importacion.plantilla("clientes")
    assert respuesta is not None
    # La respuesta es un StreamingResponse: se lee el cuerpo.
    import asyncio

    async def _leer():
        trozos = []
        async for t in respuesta.body_iterator:
            trozos.append(t)
        return b"".join(trozos)

    contenido = asyncio.run(_leer())
    wb = load_workbook(io.BytesIO(contenido))
    encabezados = [c.value for c in wb.active[1]]
    assert encabezados[:2] == ["Nombre", "Planta"]
    assert "Observaciones" in encabezados

    # Y la plantilla se puede volver a leer sin errores.
    filas, error = importacion.leer_planilla(contenido, "clientes")
    assert error is None
    assert len(filas) == 1


# ── El flujo HTTP ──────────────────────────────────────────────────────────
def test_la_pagina_de_importar_pide_autenticacion(client):
    r = client.get("/importar", follow_redirects=False)
    assert r.status_code == 303


def test_la_pagina_de_importar_carga(client, seed_admin):
    client.post("/auth/login", data={"username": "admin", "password": "admin123"})
    r = client.get("/importar")
    assert r.status_code == 200
    assert "Importar desde Excel" in r.text
    assert "Descargar plantilla" in r.text


def test_la_previsualizacion_no_escribe_en_la_base(client, db_session, seed_admin):
    """Ver la previsualización no debe crear nada."""
    client.post("/auth/login", data={"username": "admin", "password": "admin123"})
    contenido = _planilla(ENCABEZADOS_CLIENTES, [
        ["Cartocor SA", "P1", "", "", "", "", "", "", "", "", ""],
    ])
    r = client.post("/importar/previsualizar",
                    data={"tipo": "clientes", "csrf_token": _csrf(client)},
                    files={"archivo": ("c.xlsx", contenido,
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200
    assert "Cartocor SA" in r.text
    assert "Importar 1 registro" in r.text
    # Lo importante: todavía no se escribió nada.
    assert db_session.query(models.Customer).filter_by(name="Cartocor SA").count() == 0


def test_el_flujo_completo_importa(client, db_session, seed_admin):
    """Subir, previsualizar y confirmar deja el registro en la base."""
    import re

    client.post("/auth/login", data={"username": "admin", "password": "admin123"})
    contenido = _planilla(ENCABEZADOS_CLIENTES, [
        ["Cartocor SA", "P1", "Corrugado", "", "", "", "", "", "", "Activo", ""],
    ])
    r = client.post("/importar/previsualizar",
                    data={"tipo": "clientes", "csrf_token": _csrf(client)},
                    files={"archivo": ("c.xlsx", contenido, "application/octet-stream")})
    assert r.status_code == 200

    token = re.search(r'name="token" value="([^"]+)"', r.text)
    assert token, "la previsualización no devolvió el token"

    r2 = client.post("/importar/confirmar",
                     data={"tipo": "clientes", "token": token.group(1),
                           "csrf_token": _csrf(client)},
                     follow_redirects=False)
    assert r2.status_code == 303
    assert db_session.query(models.Customer).filter_by(name="Cartocor SA").count() == 1


def test_un_token_inventado_no_rompe(client, seed_admin):
    """Un token que no existe simplemente redirige, sin error 500."""
    client.post("/auth/login", data={"username": "admin", "password": "admin123"})
    r = client.post("/importar/confirmar",
                    data={"tipo": "clientes", "token": "no-es-un-uuid",
                          "csrf_token": _csrf(client)},
                    follow_redirects=False)
    assert r.status_code == 303


def test_un_token_con_ruta_no_escapa_del_directorio(client, db_session, seed_admin):
    """El token se valida como UUID: no se puede pedir otro archivo del disco."""
    client.post("/auth/login", data={"username": "admin", "password": "admin123"})
    r = client.post("/importar/confirmar",
                    data={"tipo": "clientes", "token": "../../../etc/passwd",
                          "csrf_token": _csrf(client)},
                    follow_redirects=False)
    assert r.status_code == 303
    assert db_session.query(models.Customer).count() == 0


# ── La columna Equipo ──────────────────────────────────────────────────────
#
# A partir de los equipos, un cliente se asigna a uno. La planilla puede traer
# la columna «Equipo», que manda cuando está; las tres columnas de personas se
# mantienen para que las planillas que ya se usan sigan funcionando.

def _fila_cliente(**extra):
    datos = {"_fila": 2, "name": "CARTOCOR", "plant": "Planta 1",
             "assigned_commercial": "", "designer": "", "technical_advisor": "",
             "equipo": ""}
    datos.update(extra)
    return datos


def test_la_columna_equipo_esta_en_la_plantilla():
    etiquetas = [c[1] for c in importacion.CLIENTES["columnas"]]
    assert "Equipo" in etiquetas
    # Y el ejemplo tiene un valor por columna.
    assert len(importacion.CLIENTES["ejemplo"]) == len(etiquetas)


def test_un_equipo_que_existe_se_usa_y_manda_sus_miembros(db_session):
    from services import team_service

    equipo = team_service.crear(db_session, "Equipo Norte", "Pérez", "Gómez", "López")
    listas, errores, avisos = importacion.validar(db_session, "clientes", [
        _fila_cliente(equipo=equipo.name)])

    assert not errores
    assert not avisos
    assert listas[0]["team_id"] == equipo.id
    # Los tres nombres salen del equipo, no de la planilla.
    assert listas[0]["assigned_commercial"] == equipo.commercial
    assert listas[0]["technical_advisor"] == equipo.technical_advisor
    assert listas[0]["designer"] == equipo.designer


def test_un_equipo_inexistente_con_los_tres_miembros_se_crea(db_session):
    listas, errores, avisos = importacion.validar(db_session, "clientes", [
        _fila_cliente(equipo="Equipo Nuevo", assigned_commercial="Pérez",
                      designer="Gómez", technical_advisor="López")])

    assert not errores
    assert listas[0]["team_id"] is not None
    # Se creó, y con otro nombre: hay que avisarlo.
    assert len(avisos) == 1
    assert "no existe" in avisos[0]["motivo"]


def test_un_equipo_inexistente_sin_los_miembros_rechaza_la_fila(db_session):
    """No se inventa un equipo a medias."""
    listas, errores, avisos = importacion.validar(db_session, "clientes", [
        _fila_cliente(equipo="Equipo Fantasma", assigned_commercial="Pérez")])

    assert listas == []
    assert len(errores) == 1
    assert "no existe" in errores[0]["motivo"]


def test_si_los_miembros_no_coinciden_con_el_equipo_se_avisa(db_session):
    """La columna Equipo manda, y la fila se importa igual: por eso es aviso y no error."""
    from services import team_service

    equipo = team_service.crear(db_session, "Equipo Norte", "Pérez", "Gómez", "López")
    listas, errores, avisos = importacion.validar(db_session, "clientes", [
        _fila_cliente(equipo=equipo.name, assigned_commercial="OTRO",
                      designer="OTRO2", technical_advisor="OTRO3")])

    assert not errores, "la fila tiene que entrar"
    assert len(listas) == 1
    assert len(avisos) == 1
    assert "Se usó el equipo" in avisos[0]["motivo"]
    assert listas[0]["assigned_commercial"] == equipo.commercial


def test_sin_columna_equipo_se_resuelve_por_los_tres_nombres(db_session):
    """El camino viejo sigue funcionando: es lo que usan las planillas de siempre."""
    listas, errores, avisos = importacion.validar(db_session, "clientes", [
        _fila_cliente(assigned_commercial="Pérez", designer="Gómez",
                      technical_advisor="López")])

    assert not errores
    assert listas[0]["team_id"] is not None
    # Se creó sin aviso: no había ningún nombre de equipo que contradijera nada.
    assert not avisos


def test_la_previsualizacion_avisa_sin_rechazar(client, db_session, seed_admin):
    """Una fila que entra con una advertencia tiene que verse distinto de una que no entra.

    Se sube una planilla que nombra un equipo inexistente pero trae los tres
    miembros: la fila se importa y el equipo se crea al confirmar, con un aviso. Si
    el aviso cayera en el bloque de errores —cuyo texto dice «no se van a
    importar»—, el usuario creería que la fila se descartó.
    """
    import re

    client.post("/auth/login", data={"username": "admin", "password": "admin123"})
    contenido = _planilla(ENCABEZADOS_CLIENTES, [
        ["Cartocor SA", "P1", "Corrugado", "Equipo Inventado", "Pérez", "ING. GOMEZ",
         "Diseño", "Calle 1", "Rosario", "Activo", ""],
    ])
    r = client.post("/importar/previsualizar",
                    data={"tipo": "clientes", "csrf_token": _csrf(client)},
                    files={"archivo": ("c.xlsx", contenido, "application/octet-stream")})

    assert r.status_code == 200
    assert "con una advertencia" in r.text
    assert "Importar 1 registro" in r.text
    # El aviso nombra el equipo que se va a crear.
    assert "se va a crear como" in r.text

    # La previsualización no escribe nada: ni el cliente ni el equipo.
    assert db_session.query(models.Customer).filter_by(name="Cartocor SA").count() == 0
    assert db_session.query(models.Team).count() == 0

    # Al confirmar sí se crean los dos.
    token = re.search(r'name="token" value="([^"]+)"', r.text).group(1)
    client.post("/importar/confirmar",
                data={"tipo": "clientes", "token": token,
                      "csrf_token": _csrf(client)},
                follow_redirects=False)

    assert db_session.query(models.Customer).filter_by(name="Cartocor SA").count() == 1
    equipo = db_session.query(models.Team).filter_by(
        name="PÉREZ / DISEÑO / ING. GOMEZ").first()
    assert equipo is not None, "no se creó el equipo a partir de los tres nombres"
    cliente = db_session.query(models.Customer).filter_by(name="Cartocor SA").one()
    assert cliente.team_id == equipo.id
