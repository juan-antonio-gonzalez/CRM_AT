"""Pruebas de los filtros y la exportación del listado de minutas.

El listado sólo tenía paginación: con muchas actas no había forma de encontrar
una reunión por su objetivo, su tipo, quién la redactó o el rango de fechas.
"""

import io
import re

from openpyxl import load_workbook

import models


def _seed(db, cantidad=8):
    """Crea minutas alternando estado, tipo y redactor."""
    db.query(models.MeetingMinutes).delete()
    db.commit()
    for i in range(cantidad):
        finalizada = i % 2 == 0
        db.add(models.MeetingMinutes(
            fecha=f"2026-0{(i % 6) + 1}-1{i % 9}",
            status="Finalizado" if finalizada else "Borrador",
            form_data={
                "fecha": f"2026-0{(i % 6) + 1}-1{i % 9}",
                "horaInicio": "09:00", "horaFin": "10:30",
                "lugar": "Sala de reuniones" if i % 2 == 0 else "Virtual",
                "tipoReunion": "Seguimiento" if i < cantidad // 2 else "Comercial",
                "objetivo": f"Objetivo de la reunión número {i:02d}",
                "redactor": "Ing. Gómez" if i % 2 == 0 else "Ing. Torres",
                "presentes": [{"nombre": "Ana"}, {"nombre": "Luis"}],
                "ausentes": [],
                "acuerdos": [{"acuerdo": "Acuerdo pendiente de aprobación",
                              "observaciones": "Con seguimiento mensual"}],
                "tareas": [{"tarea": "Enviar informe", "responsable": "Ana",
                            "plazo": "2026-07-01", "estado": "Pendiente"}],
            },
        ))
    db.commit()


def _filas(html):
    return html.count('<tr class="hover:bg-slate-50')


def test_sin_filtros_lista_todo(client, db_session, admin_cookie):
    _seed(db_session, 8)
    html = client.get("/minutas").text
    assert _filas(html) == 8
    assert "8 minutas" in html


def test_filtro_por_estado(client, db_session, admin_cookie):
    _seed(db_session, 8)
    assert _filas(client.get("/minutas?estado=Finalizado").text) == 4
    assert _filas(client.get("/minutas?estado=Borrador").text) == 4


def test_filtro_por_tipo_de_reunion(client, db_session, admin_cookie):
    _seed(db_session, 8)
    html = client.get("/minutas?tipo=Seguimiento").text
    assert _filas(html) == 4
    assert "Comercial" not in html.split("<tbody")[1].split("</tbody>")[0]


def test_filtro_por_redactor(client, db_session, admin_cookie):
    _seed(db_session, 8)
    assert _filas(client.get("/minutas?redactor=Torres").text) == 4


def test_busqueda_por_objetivo_y_lugar(client, db_session, admin_cookie):
    _seed(db_session, 8)
    assert _filas(client.get("/minutas?q=número+03").text) == 1
    assert _filas(client.get("/minutas?q=Virtual").text) == 4
    assert _filas(client.get("/minutas?q=no-existe-este-texto").text) == 0


def test_filtro_por_rango_de_fechas(client, db_session, admin_cookie):
    _seed(db_session, 8)
    todos = _filas(client.get("/minutas").text)
    pocos = _filas(client.get("/minutas?fecha_desde=2026-03-01&fecha_hasta=2026-04-30").text)
    assert 0 <= pocos < todos


def test_paginacion_conserva_los_filtros(client, db_session, admin_cookie):
    _seed(db_session, 30)
    # 30 registros alternan estado, así que hay 15 borradores: alcanza para 1
    # página de 20 y no para 2. Se usa un per_page chico para forzar 2 páginas.
    html = client.get("/minutas?estado=Borrador&per_page=10&page=2").text
    assert "Página 2" in html
    # los enlaces de paginación deben arrastrar el filtro
    assert "estado=Borrador" in html


def test_exportar_minutas_a_excel(client, db_session, admin_cookie):
    _seed(db_session, 5)
    r = client.get("/minutas/exportar-excel")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert r.content[:2] == b"PK"

    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.cell(row=1, column=1).value == "Nº"
    assert ws.max_row == 6, "5 filas de datos más el encabezado"
    # los campos del JSON deben salir aplanados y legibles
    encabezados = [c.value for c in ws[1]]
    assert "Acuerdos" in encabezados and "Presentes" in encabezados
    fila = [ws.cell(row=2, column=i).value for i in range(1, 14)]  # 13 columnas
    assert any("Ana" in str(v) for v in fila if v), "los presentes no se exportaron"
    assert any("Acuerdo pendiente" in str(v) for v in fila if v), "los acuerdos no se exportaron"


def test_exportar_respeta_el_filtro(client, db_session, admin_cookie):
    _seed(db_session, 8)
    r = client.get("/minutas/exportar-excel?estado=Finalizado")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == 5, "4 finalizadas más el encabezado"


def test_estado_vacio_distingue_filtro_de_lista_vacia(client, db_session, admin_cookie):
    db_session.query(models.MeetingMinutes).delete()
    db_session.commit()

    html = client.get("/minutas").text
    assert "No hay minutas" in html

    html = client.get("/minutas?q=nadie").text
    assert "Sin resultados" in html
    assert "Limpiar filtros" in html
