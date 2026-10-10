"""Pruebas del vínculo visita técnica <-> plan de acción, y del CRUD de seguimiento.

Estas pruebas existen porque ese vínculo tenía dos fallos reales:

* al eliminar una visita, sus tareas del plan de acción quedaban **activas y
  huérfanas** (seguían apareciendo en /seguimiento-plan apuntando a una visita
  inexistente), aunque el aviso de borrado prometía eliminarlas;
* al restaurar la visita desde la papelera se creaba con un **id nuevo**, así que
  las tareas seguían apuntando al id viejo y el vínculo quedaba roto.

También cubren el CRUD del módulo de seguimiento, que no tenía forma de crear
una tarea sin pasar por una visita.
"""

import json
import re

import pytest

import models


def _csrf(client, ruta="/visitas2"):
    """El token viaja en una meta etiqueta y app.js lo inyecta en los formularios."""
    html = client.get(ruta).text
    m = re.search(r'name="csrf-token" content="([^"]+)"', html)
    assert m, f"no se encontró el token CSRF en {ruta}"
    return m.group(1)


def _payload(tarea, estado_visita="Finalizado", fecha="2026-05-10"):
    return {
        "fecha": fecha,
        "status": estado_visita,
        "cliente_nombre": "ACME SA (Planta Norte)",
        "at": "ING. GOMEZ",
        "objetivo": "Auditoría anual",
        "tables": {"9": [["ET-001", tarea, "Roberto", "15/06/2026"]]},
    }


def _crear_visita(client, csrf, tarea, estado_visita="Finalizado", fecha="2026-05-10"):
    r = client.post("/visitas2/nueva",
                    data={"form_data": json.dumps(_payload(tarea, estado_visita, fecha)),
                          "csrf_token": csrf},
                    follow_redirects=False)
    assert r.status_code == 303
    return r


def _editar_visita(client, csrf, visit_id, tarea, tabla9=None):
    """Reguarda una visita existente. ``tabla9`` permite vaciar las acciones."""
    data = _payload(tarea)
    if tabla9 is not None:
        data["tables"]["9"] = tabla9
    r = client.post(f"/visitas2/{visit_id}/editar",
                    data={"form_data": json.dumps(data), "csrf_token": csrf},
                    follow_redirects=False)
    assert r.status_code == 303
    return r


def _tareas(db):
    return db.query(models.ActionPlanTracking).order_by(models.ActionPlanTracking.id).all()


# ── Vínculo visita <-> plan de acción ──────────────────────────────────────
def test_crear_visita_vincula_las_acciones(client, db_session, admin_cookie):
    """La tabla 9 del formulario debe crear tareas ligadas a la visita."""
    csrf = _csrf(client)
    _crear_visita(client, csrf, "Reemplazar rodamiento")

    visita = db_session.query(models.TechnicalVisit2).one()
    tareas = _tareas(db_session)

    assert len(tareas) == 1
    assert tareas[0].task == "Reemplazar rodamiento"
    assert tareas[0].source_visit_id == visita.id
    assert tareas[0].is_active is True
    assert tareas[0].status == "Pendiente"


def test_editar_visita_no_pierde_el_avance_de_las_tareas(client, db_session, admin_cookie):
    """Reguardar la visita no debe perder el estado ni duplicar las tareas."""
    csrf = _csrf(client)
    _crear_visita(client, csrf, "Reemplazar rodamiento")
    visita = db_session.query(models.TechnicalVisit2).one()

    tarea = _tareas(db_session)[0]
    tarea.status = "Finalizado"
    tarea.observations = "Avance del usuario"
    db_session.commit()

    # Reguardar la MISMA visita, no crear otra.
    _editar_visita(client, csrf, visita.id, "Reemplazar rodamiento")

    assert db_session.query(models.TechnicalVisit2).count() == 1, "se creó una visita de más"
    tareas = _tareas(db_session)
    assert len(tareas) == 1, "se duplicaron las tareas al reguardar la visita"
    assert tareas[0].status == "Finalizado", "se perdió el avance registrado"
    assert tareas[0].observations == "Avance del usuario"


def test_editar_visita_desactiva_las_tareas_eliminadas(client, db_session, admin_cookie):
    """Si se quita una acción del formulario, su tarea se desactiva (no se borra)."""
    csrf = _csrf(client)
    _crear_visita(client, csrf, "Reemplazar rodamiento")
    visita = db_session.query(models.TechnicalVisit2).one()

    _editar_visita(client, csrf, visita.id, "Reemplazar rodamiento", tabla9=[])

    tareas = _tareas(db_session)
    assert len(tareas) == 1, "la tarea no debe borrarse, sólo desactivarse"
    assert tareas[0].is_active is False


def test_eliminar_visita_desactiva_sus_tareas(client, db_session, admin_cookie):
    """Eliminar la visita no debe dejar tareas huérfanas activas."""
    csrf = _csrf(client)
    _crear_visita(client, csrf, "Reemplazar rodamiento")
    visita = db_session.query(models.TechnicalVisit2).one()

    client.post(f"/eliminar-visita2/{visita.id}", data={"csrf_token": csrf},
                follow_redirects=False)

    assert db_session.query(models.TechnicalVisit2).count() == 0
    tareas = _tareas(db_session)
    assert len(tareas) == 1, "la tarea se conserva para poder restaurarla"
    assert tareas[0].is_active is False, "quedó una tarea huérfana activa"


def test_restaurar_visita_revincula_sus_tareas(client, db_session, admin_cookie):
    """Al restaurar, las tareas deben apuntar al id nuevo y reactivarse."""
    csrf = _csrf(client)
    _crear_visita(client, csrf, "Reemplazar rodamiento")
    visita_id = db_session.query(models.TechnicalVisit2).one().id

    client.post(f"/eliminar-visita2/{visita_id}", data={"csrf_token": csrf},
                follow_redirects=False)

    item = db_session.query(models.TrashItem).filter_by(entity_type="visit2").one()
    client.post(f"/papelera/restaurar/{item.id}", data={"csrf_token": csrf},
                follow_redirects=False)

    restaurada = db_session.query(models.TechnicalVisit2).one()
    assert restaurada.id != visita_id, "la restauración crea un id nuevo"

    tarea = _tareas(db_session)[0]
    assert tarea.source_visit_id == restaurada.id, "el vínculo quedó apuntando al id viejo"
    assert tarea.is_active is True, "la tarea no se reactivó"


# ── CRUD del módulo de seguimiento ─────────────────────────────────────────
def test_crear_tarea_de_seguimiento_manualmente(client, db_session, admin_cookie):
    """El módulo debe permitir crear una tarea fuera del flujo de visitas."""
    csrf = _csrf(client, "/seguimiento-plan")
    r = client.post("/seguimiento-plan/nueva", data={
        "csrf_token": csrf,
        "task": "Enviar cotización de repuestos",
        "customer_name": "ACME SA",
        "et_number": "ET-777",
        "responsible": "Roberto",
        "due_date": "30/06/2026",
        "status": "Pendiente",
        "observations": "Surgió de una llamada",
    }, follow_redirects=False)

    assert r.status_code == 303
    tareas = _tareas(db_session)
    assert len(tareas) == 1
    assert tareas[0].task == "Enviar cotización de repuestos"
    assert tareas[0].source_visit_id is None, "no nace de una visita"
    assert tareas[0].status == "Pendiente"
    assert tareas[0].is_active is True


def test_crear_tarea_sin_texto_se_rechaza(client, db_session, admin_cookie):
    csrf = _csrf(client, "/seguimiento-plan")
    client.post("/seguimiento-plan/nueva",
                data={"csrf_token": csrf, "task": "   "}, follow_redirects=False)
    assert _tareas(db_session) == []


def test_editar_y_eliminar_tarea(client, db_session, admin_cookie):
    csrf = _csrf(client, "/seguimiento-plan")
    client.post("/seguimiento-plan/nueva",
                data={"csrf_token": csrf, "task": "Tarea X", "status": "Pendiente"},
                follow_redirects=False)
    tarea = _tareas(db_session)[0]

    client.post(f"/seguimiento-plan/{tarea.id}/editar", data={
        "csrf_token": csrf, "task": "Tarea X", "status": "En curso",
        "observations": "En progreso",
    }, follow_redirects=False)
    db_session.refresh(tarea)
    assert tarea.status == "En curso"
    assert tarea.observations == "En progreso"

    client.post(f"/seguimiento-plan/{tarea.id}/eliminar",
                data={"csrf_token": csrf}, follow_redirects=False)
    db_session.refresh(tarea)
    assert tarea.is_active is False, "debe ser baja lógica, no borrado"


def test_buscar_tareas_por_texto_y_estado(client, db_session, admin_cookie):
    csrf = _csrf(client, "/seguimiento-plan")
    for tarea, estado in [("Cambiar filtro de aceite", "Pendiente"),
                          ("Calibrar sensor", "Finalizado")]:
        client.post("/seguimiento-plan/nueva",
                    data={"csrf_token": csrf, "task": tarea, "status": estado},
                    follow_redirects=False)

    assert "Cambiar filtro de aceite" in client.get("/seguimiento-plan?q=aceite").text
    assert "Calibrar sensor" not in client.get("/seguimiento-plan?q=aceite").text
    assert "Calibrar sensor" in client.get("/seguimiento-plan?estado=Finalizado").text
    assert "Cambiar filtro" not in client.get("/seguimiento-plan?estado=Finalizado").text


# ── Filtros y exportación de visitas ───────────────────────────────────────
def test_filtros_de_visitas(client, db_session, admin_cookie):
    csrf = _csrf(client)
    _crear_visita(client, csrf, "Tarea A", estado_visita="Finalizado", fecha="2026-01-10")
    _crear_visita(client, csrf, "Tarea B", estado_visita="Programado", fecha="2026-03-20")

    def filas(**params):
        qs = "&".join(f"{k}={v}" for k, v in params.items())
        return client.get(f"/visitas2?{qs}").text.count('<tr class="hover:bg-slate-50')

    assert filas() == 2
    assert filas(estado="Finalizado") == 1
    assert filas(estado="Programado") == 1
    assert filas(fecha_desde="2026-02-01") == 1
    assert filas(fecha_hasta="2026-02-01") == 1
    assert filas(q="ACME") == 2
    assert filas(q="no-existe-este-texto") == 0


def test_exportar_visitas_a_excel(client, db_session, admin_cookie):
    import io
    from openpyxl import load_workbook

    csrf = _csrf(client)
    _crear_visita(client, csrf, "Tarea A", estado_visita="Finalizado")

    r = client.get("/visitas2/exportar-excel")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert r.content[:2] == b"PK", "no parece un xlsx"

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert ws.max_row == 2, "debería haber 1 fila de datos más el encabezado"
    assert ws.cell(row=1, column=1).value == "Fecha"
    assert ws.cell(row=2, column=2).value == "ACME SA (Planta Norte)"


def test_enlaces_de_navegacion_no_apuntan_a_rutas_inexistentes(client, db_session, admin_cookie):
    """El botón de volver apuntaba a /dashboard, que devuelve 404."""
    html = client.get("/visitas2").text
    assert 'href="/dashboard"' not in html
    assert client.get("/dashboard").status_code == 404


def test_el_formulario_de_visita_muestra_sus_tareas_vinculadas(client, db_session, admin_cookie):
    """El vínculo debe ser visible en el formulario, no sólo en /seguimiento-plan."""
    csrf = _csrf(client)
    _crear_visita(client, csrf, "Reemplazar rodamiento")
    visita = db_session.query(models.TechnicalVisit2).one()

    html = client.get(f"/visitas2/{visita.id}/editar").text
    assert "Tareas del plan de acción generadas por esta visita" in html
    assert "Reemplazar rodamiento" in html
    assert "Seguimiento Plan de Acción" in html

    # Al avanzar la tarea en seguimiento, el formulario refleja el estado nuevo.
    tarea = _tareas(db_session)[0]
    tarea.status = "En curso"
    db_session.commit()
    html = client.get(f"/visitas2/{visita.id}/editar").text
    assert "En curso" in html


def test_visita_nueva_no_muestra_tareas_vinculadas(client, db_session, admin_cookie):
    """Al crear, todavía no hay tareas que mostrar."""
    html = client.get("/visitas2/nueva").text
    assert "Tareas del plan de acción generadas por esta visita" not in html
