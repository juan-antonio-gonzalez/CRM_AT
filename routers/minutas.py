from fastapi import APIRouter, Request, Depends, Query
from fastapi.responses import HTMLResponse, RedirectResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, or_
import io
import models
from database import get_db
from services import pagination as pg
from services import auditoria
from services.utils import safe_json_loads
from services.customer_service import list_customers
import datetime
from templating import templates

router = APIRouter(tags=["minutas"])

# Estados que escribe el formulario de minuta.
MINUTA_STATUSES = ["Borrador", "Finalizado"]


def _minuta_filters(db: Session, q: str = None, estado: str = None,
                    tipo: str = None, fecha_desde: str = None,
                    fecha_hasta: str = None, redactor: str = None):
    """Consulta base del listado de minutas, con todos los filtros aplicados.

    El objetivo, el lugar, el tipo de reunión y el redactor viven dentro del JSON
    del formulario, así que se filtran sobre esos campos.
    """
    query = db.query(models.MeetingMinutes)

    if estado and estado.strip():
        query = query.filter(models.MeetingMinutes.status == estado.strip())

    if tipo and tipo.strip():
        query = query.filter(
            models.MeetingMinutes.form_data["tipoReunion"].as_string() == tipo.strip()
        )

    if redactor and redactor.strip():
        query = query.filter(
            models.MeetingMinutes.form_data["redactor"].as_string().ilike(f"%{redactor.strip()}%")
        )

    if fecha_desde:
        query = query.filter(models.MeetingMinutes.fecha >= fecha_desde)
    if fecha_hasta:
        query = query.filter(models.MeetingMinutes.fecha <= fecha_hasta)

    if q and q.strip():
        patron = f"%{q.strip()}%"
        query = query.filter(or_(
            models.MeetingMinutes.fecha.ilike(patron),
            models.MeetingMinutes.status.ilike(patron),
            models.MeetingMinutes.form_data["objetivo"].as_string().ilike(patron),
            models.MeetingMinutes.form_data["lugar"].as_string().ilike(patron),
            models.MeetingMinutes.form_data["tipoReunion"].as_string().ilike(patron),
            models.MeetingMinutes.form_data["redactor"].as_string().ilike(patron),
        ))

    return query


def _query_string(**kwargs) -> str:
    """Query string con los filtros que tienen valor, para paginar sin perderlos."""
    from urllib.parse import urlencode
    return urlencode({k: v for k, v in kwargs.items() if v not in (None, "", "Todos")})


@router.get("/minutas", response_class=HTMLResponse)
async def minutas_lista(
    request: Request,
    page: int = Query(1, ge=1),
    per_page: int = Query(20, ge=1, le=100),
    q: str = Query(None),
    estado: str = Query(None),
    tipo: str = Query(None),
    redactor: str = Query(None),
    fecha_desde: str = Query(None),
    fecha_hasta: str = Query(None),
    db: Session = Depends(get_db)
):
    filtros = dict(q=q, estado=estado, tipo=tipo, redactor=redactor,
                   fecha_desde=fecha_desde, fecha_hasta=fecha_hasta)

    base = _minuta_filters(db, **filtros)
    total_filtrado = base.count()

    page, per_page, offset = pg.paginate(page, per_page)
    minutas = base.order_by(models.MeetingMinutes.id.desc()).offset(offset).limit(per_page).all()

    # Las tarjetas describen el conjunto completo, en una sola consulta agrupada.
    conteos = dict(
        db.query(models.MeetingMinutes.status, func.count(models.MeetingMinutes.id))
        .group_by(models.MeetingMinutes.status)
        .all()
    )
    total = sum(conteos.values())
    stats = {
        "total": total,
        "finalizadas": conteos.get("Finalizado", 0),
        "borradores": conteos.get("Borrador", 0),
    }

    tipos = sorted({
        t[0] for t in db.query(models.MeetingMinutes.form_data["tipoReunion"].as_string())
        .distinct().all() if t[0]
    })

    return templates.TemplateResponse(request, 'minutas_lista.html', {
        "request": request,
        "minutas": minutas,
        "pagination": pg.build_pagination(page, per_page, total_filtrado),
        "filtered_total": total_filtrado,
        "stats": stats,
        "active_page": "minutas",
        "estados": MINUTA_STATUSES,
        "tipos": tipos,
        "filtros": filtros,
        "qs": _query_string(**filtros),
    })


@router.get("/minutas/exportar-excel")
async def exportar_minutas_excel(
    q: str = Query(None),
    estado: str = Query(None),
    tipo: str = Query(None),
    redactor: str = Query(None),
    fecha_desde: str = Query(None),
    fecha_hasta: str = Query(None),
    db: Session = Depends(get_db),
):
    """Exporta el listado filtrado de minutas a Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from fastapi.responses import StreamingResponse

    minutas = (
        _minuta_filters(db, q=q, estado=estado, tipo=tipo, redactor=redactor,
                        fecha_desde=fecha_desde, fecha_hasta=fecha_hasta)
        .order_by(models.MeetingMinutes.id.desc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Minutas"
    ws.append(["Nº", "Fecha", "Hora inicio", "Hora fin", "Lugar / modalidad",
               "Tipo de reunión", "Objetivo", "Redactor", "Estado",
               "Presentes", "Ausentes", "Acuerdos", "Tareas"])
    for celda in ws[1]:
        celda.font = Font(bold=True)

    for m in minutas:
        fd = m.form_data or {}
        presentes = ", ".join(p.get("nombre", "") for p in (fd.get("presentes") or []))
        ausentes = ", ".join(p.get("nombre", "") for p in (fd.get("ausentes") or []))
        acuerdos = " | ".join(a.get("acuerdo", "") for a in (fd.get("acuerdos") or []))
        tareas = " | ".join(
            f"{t.get('tarea', '')} ({t.get('responsable', '')} - {t.get('plazo', '')})"
            for t in (fd.get("tareas") or [])
        )
        ws.append([
            m.id, m.fecha or "", fd.get("horaInicio", ""), fd.get("horaFin", ""),
            fd.get("lugar", ""), fd.get("tipoReunion", ""), fd.get("objetivo", ""),
            fd.get("redactor", ""), m.status or "",
            presentes, ausentes, acuerdos, tareas,
        ])

    for i, ancho in enumerate([6, 12, 11, 10, 22, 18, 40, 22, 13, 30, 20, 50, 50], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    nombre = f"minutas-{datetime.date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre}"},
    )


@router.get("/minutas/nueva", response_class=HTMLResponse)
async def nueva_minuta(request: Request, db: Session = Depends(get_db)):
    clientes = list_customers(db)
    return templates.TemplateResponse(request, 'nueva_minuta.html', {
        "request": request,
        "minuta": None,
        "form_data": None,
        "title": "Nueva Minuta de Reunión",
        "action_url": "/minutas/nueva",
        "active_page": "minutas",
        "clientes": clientes,
    })


def _etiqueta_minuta(minuta):
    """Cómo se nombra una minuta en el historial."""
    if minuta is None:
        return None
    fd = minuta.form_data or {}
    objetivo = fd.get("objetivo") or ""
    if objetivo:
        return objetivo[:200]
    return f"minuta del {minuta.fecha}" if minuta.fecha else f"minuta #{minuta.id}"


@router.post("/minutas/nueva")
async def create_minuta(request: Request, db: Session = Depends(get_db)):
    form = await request.form()
    data = safe_json_loads(form.get("form_data"))

    fecha = data.get("fecha", "") or ""
    estado = data.get("status", "Borrador")
    new_minuta = models.MeetingMinutes(
        fecha=fecha,
        status=estado,
        form_data=data,
    )
    db.add(new_minuta)
    db.commit()
    db.refresh(new_minuta)
    auditoria.registrar(db, "minute", new_minuta.id, "crear",
                        user=getattr(request.state, "user", None),
                        entity_label=_etiqueta_minuta(new_minuta))

    return RedirectResponse(url="/minutas", status_code=303)


@router.get("/minutas/{minuta_id}/editar", response_class=HTMLResponse)
async def editar_minuta(request: Request, minuta_id: int, db: Session = Depends(get_db)):
    minuta = db.query(models.MeetingMinutes).filter(models.MeetingMinutes.id == minuta_id).first()
    if not minuta:
        return RedirectResponse(url="/minutas", status_code=303)

    clientes = list_customers(db)
    return templates.TemplateResponse(request, 'nueva_minuta.html', {
        "request": request,
        "minuta": minuta,
        "form_data": minuta.form_data or {},
        "title": f"Editar Minuta #{minuta.id}",
        "action_url": f"/minutas/{minuta_id}/editar",
        "active_page": "minutas",
        "clientes": clientes,
    })


@router.post("/minutas/{minuta_id}/editar")
async def save_editar_minuta(minuta_id: int, request: Request, db: Session = Depends(get_db)):
    minuta = db.query(models.MeetingMinutes).filter(models.MeetingMinutes.id == minuta_id).first()
    if not minuta:
        return RedirectResponse(url="/minutas", status_code=303)

    # Foto previa: sin esto no se puede decir qué cambió.
    antes = auditoria.snapshot(minuta)
    form = await request.form()
    data = safe_json_loads(form.get("form_data"))

    minuta.fecha = data.get("fecha", "") or ""
    minuta.status = data.get("status", "Borrador")
    minuta.form_data = data
    minuta.updated_at = datetime.datetime.utcnow()
    db.commit()

    # El historial se registra al final, cuando el cambio ya está guardado.
    despues = auditoria.snapshot(minuta)
    cambios = auditoria.diferencias(antes, despues)
    if cambios:
        auditoria.registrar(db, "minute", minuta_id, "editar",
                            user=getattr(request.state, "user", None),
                            entity_label=_etiqueta_minuta(minuta), changes=cambios)

    return RedirectResponse(url="/minutas", status_code=303)


@router.post("/eliminar-minuta/{minuta_id}")
async def delete_minuta(minuta_id: int, request: Request, db: Session = Depends(get_db)):
    minuta = db.query(models.MeetingMinutes).filter(models.MeetingMinutes.id == minuta_id).first()
    if minuta:
        from services.trash_service import move_to_trash
        user = getattr(request.state, 'user', None)
        uid = int(user["sub"]) if user else None
        uname = user.get("username") if user else None
        fd = minuta.form_data or {}
        move_to_trash(db, "meeting_minutes", minuta.id, {
            "id": minuta.id,
            "fecha": minuta.fecha,
            "status": minuta.status,
            "form_data": fd,
        }, uid, uname, f"Minuta #{minuta.id}")
        etiqueta = _etiqueta_minuta(minuta)
        db.delete(minuta)
        db.commit()
        auditoria.registrar(db, "minute", minuta_id, "eliminar",
                            user=getattr(request.state, "user", None),
                            entity_label=etiqueta)
    return RedirectResponse(url="/minutas", status_code=303)
