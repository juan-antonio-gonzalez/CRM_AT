from fastapi import APIRouter, Request, Depends, Query
from fastapi.responses import HTMLResponse, RedirectResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, or_
import io
import models
from database import get_db
from services import pagination as pg
from services.utils import safe_json_loads
from services.customer_service import list_customers
import datetime

from enums import TrackingStatus
from templating import templates

router = APIRouter(tags=["visits2"])

# Estados reales que escribe el formulario de visitas. Se usan tanto para el
# filtro como para las tarjetas de resumen, así el vocabulario es uno solo.
VISIT_STATUSES = ["En proceso", "Programado", "Finalizado", "Cancelado"]


def _visit_filters(db: Session, q: str = None, estado: str = None,
                   cliente: str = None, at: str = None,
                   fecha_desde: str = None, fecha_hasta: str = None):
    """Consulta base de visitas con todos los filtros aplicados.

    Antes esta pantalla no tenía ningún filtro: con muchas visitas no había forma
    de responder "qué visitas hizo este cliente" o "qué pasó en marzo".
    """
    query = db.query(models.TechnicalVisit2)

    if estado and estado.strip():
        query = query.filter(models.TechnicalVisit2.status == estado.strip())

    if fecha_desde:
        query = query.filter(models.TechnicalVisit2.fecha >= fecha_desde)
    if fecha_hasta:
        query = query.filter(models.TechnicalVisit2.fecha <= fecha_hasta)

    if cliente and cliente.strip():
        # El cliente vive dentro del JSON del formulario, así que el filtro se
        # aplica sobre ese campo.
        query = query.filter(
            models.TechnicalVisit2.form_data["cliente_nombre"].as_string().ilike(f"%{cliente.strip()}%")
        )

    if at and at.strip():
        query = query.filter(
            models.TechnicalVisit2.form_data["at"].as_string() == at.strip()
        )

    if q and q.strip():
        patron = f"%{q.strip()}%"
        query = query.filter(or_(
            models.TechnicalVisit2.fecha.ilike(patron),
            models.TechnicalVisit2.form_data["cliente_nombre"].as_string().ilike(patron),
            models.TechnicalVisit2.form_data["at"].as_string().ilike(patron),
            models.TechnicalVisit2.form_data["objetivo"].as_string().ilike(patron),
            models.TechnicalVisit2.status.ilike(patron),
        ))

    return query


def _query_string(**kwargs) -> str:
    """Arma el query string conservando sólo los filtros con valor."""
    from urllib.parse import urlencode
    return urlencode({k: v for k, v in kwargs.items() if v not in (None, "", "Todos")})


def cliente_nombre_planta(cliente_text: str) -> tuple:
    """Extract name and plant from 'Name (Plant)' format."""
    if not cliente_text:
        return "", ""
    if "(" in cliente_text:
        name = cliente_text.split("(")[0].strip()
        plant = cliente_text.split("(")[1].rstrip(")").strip()
        return name, plant
    return cliente_text, ""


def sync_action_plan(db: Session, visit_id: int, data: dict):
    """Sync Table 9 (ACCIONES) rows into action_plan_tracking.

    Implemented as an upsert keyed on (source_visit_id, et_number, task).

    The previous version deleted every tracking row for the visit and recreated
    them with ``status=PENDIENTE``, so any progress and observations the user had
    recorded in /seguimiento-plan were silently wiped on the next save.
    """
    tables = data.get("tables", {})
    rows = tables.get("9", [])

    cliente_nombre, cliente_planta = cliente_nombre_planta(data.get("cliente_nombre", ""))

    existing = db.query(models.ActionPlanTracking).filter(
        models.ActionPlanTracking.source_visit_id == visit_id
    ).all()
    existing_by_key = {(r.et_number or "", r.task or ""): r for r in existing}
    seen_keys = set()

    for row in rows:
        if not isinstance(row, (list, tuple)) or len(row) < 4:
            continue
        et = (row[0] or "").strip()
        accion = (row[1] or "").strip()
        responsable = (row[2] or "").strip()
        plazo = (row[3] or "").strip()

        if not accion:
            continue  # skip empty action rows

        key = (et, accion)
        seen_keys.add(key)
        record = existing_by_key.get(key)

        if record:
            # Preserve status/observations owned by the seguimiento-plan screen.
            record.responsible = responsable
            record.due_date = plazo
            record.customer_name = cliente_nombre
            record.plant = cliente_planta
            record.updated_at = datetime.datetime.now()
        else:
            db.add(models.ActionPlanTracking(
                customer_name=cliente_nombre,
                plant=cliente_planta,
                visit_date=datetime.datetime.now(),
                et_number=et,
                task=accion,
                responsible=responsable,
                due_date=plazo,
                status=TrackingStatus.PENDIENTE.value,
                source_visit_id=visit_id,
                is_active=True,
            ))

    # Rows removed from the form are deactivated rather than deleted, so their
    # history and any recorded progress survive.
    for key, record in existing_by_key.items():
        if key not in seen_keys and record.is_active:
            record.is_active = False
            record.inactivation_reason = "Acción eliminada de la visita"
            record.inactivated_at = datetime.datetime.now()


@router.get("/visitas2", response_class=HTMLResponse)
async def visitas2_lista(
    request: Request,
    page: int = Query(1, ge=1),
    per_page: int = Query(20, ge=1, le=100),
    sort: str = Query("fecha_desc"),
    q: str = Query(None),
    estado: str = Query(None),
    cliente: str = Query(None),
    at: str = Query(None),
    fecha_desde: str = Query(None),
    fecha_hasta: str = Query(None),
    db: Session = Depends(get_db)
):
    filtros = dict(q=q, estado=estado, cliente=cliente, at=at,
                   fecha_desde=fecha_desde, fecha_hasta=fecha_hasta)

    base = _visit_filters(db, **filtros)
    total = base.count()

    if sort == "fecha_asc":
        order = models.TechnicalVisit2.fecha.asc()
    else:
        order = models.TechnicalVisit2.fecha.desc()
        sort = "fecha_desc"

    page, per_page, offset = pg.paginate(page, per_page)
    visits = base.order_by(order).offset(offset).limit(per_page).all()

    # Las tarjetas describen el conjunto completo, no sólo la página visible, y
    # salen de una sola consulta agrupada en lugar de cuatro COUNT separados.
    conteos = dict(
        db.query(models.TechnicalVisit2.status, func.count(models.TechnicalVisit2.id))
        .group_by(models.TechnicalVisit2.status)
        .all()
    )
    stats = {
        "hechas": conteos.get("Finalizado", 0),
        "programadas": conteos.get("Programado", 0),
        "en_proceso": conteos.get("En proceso", 0),
        "canceladas": conteos.get("Cancelado", 0),
    }

    advisors = [
        a[0] for a in db.query(models.TechnicalVisit2.form_data["at"].as_string())
        .distinct().all() if a[0]
    ]

    return templates.TemplateResponse(request, 'visitas2_lista.html', {
        "request": request,
        "visits": visits,
        "pagination": pg.build_pagination(page, per_page, total),
        "filtered_total": total,
        "stats": stats,
        "active_page": "visitas2",
        "sort": sort,
        "estados": VISIT_STATUSES,
        "advisors": sorted(advisors),
        "filtros": filtros,
        "qs": _query_string(sort=sort, **filtros),
    })


@router.get("/visitas2/exportar-excel")
async def exportar_visitas_excel(
    request: Request,
    q: str = Query(None),
    estado: str = Query(None),
    cliente: str = Query(None),
    at: str = Query(None),
    fecha_desde: str = Query(None),
    fecha_hasta: str = Query(None),
    db: Session = Depends(get_db),
):
    """Exporta el listado filtrado a Excel (mismo patrón que /clientes)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from fastapi.responses import StreamingResponse

    visits = (
        _visit_filters(db, q=q, estado=estado, cliente=cliente, at=at,
                       fecha_desde=fecha_desde, fecha_hasta=fecha_hasta)
        .order_by(models.TechnicalVisit2.fecha.desc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Visitas técnicas"
    encabezados = ["Fecha", "Cliente", "Planta", "AT", "Estado", "Objetivo",
                   "Tipo de producto", "Planta (distribuidora)"]
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True)

    for v in visits:
        fd = v.form_data or {}
        ws.append([
            v.fecha or "",
            fd.get("cliente_nombre", ""),
            fd.get("planta", ""),
            fd.get("at", ""),
            v.status or "",
            fd.get("objetivo", ""),
            fd.get("tipo_producto", ""),
            fd.get("planta_distribuidora", ""),
        ])

    anchos = [12, 30, 18, 20, 14, 30, 20, 22]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    nombre = f"visitas-{datetime.date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre}"},
    )


@router.get("/visitas2/nueva", response_class=HTMLResponse)
async def nueva_visita2(request: Request, db: Session = Depends(get_db)):
    clientes = list_customers(db)
    tecnicos = db.query(models.TechnicalAdvisor).all()
    distribuidoras = db.query(models.DistributorPlant).order_by(models.DistributorPlant.name).all()
    temas_calidad = db.query(models.QualityTopic).order_by(models.QualityTopic.name).all()
    calidad_at = db.query(models.QualityAT).order_by(models.QualityAT.name).all()
    return templates.TemplateResponse(request, 'nueva_visita2.html', {
        "request": request,
        "visit": None,
        "form_data": None,
        "title": "Nueva Visita Técnica V2",
        "action_url": "/visitas2/nueva",
        "active_page": "visitas2",
        "clientes": clientes,
        "tecnicos": tecnicos,
        "distribuidoras": distribuidoras,
        "temas_calidad": temas_calidad,
        "calidad_at": calidad_at
    })


@router.post("/visitas2/nueva")
async def create_visita2(request: Request, db: Session = Depends(get_db)):
    form = await request.form()
    data = safe_json_loads(form.get("form_data"))

    fecha = data.get("fecha", "") or ""
    estado = data.get("status", "En proceso")
    new_visit = models.TechnicalVisit2(
        fecha=fecha,
        status=estado,
        form_data=data,
    )
    db.add(new_visit)
    db.commit()
    db.refresh(new_visit)

    sync_action_plan(db, new_visit.id, data)
    db.commit()

    return RedirectResponse(url="/visitas2", status_code=303)


@router.get("/visitas2/{visit_id}/editar", response_class=HTMLResponse)
async def editar_visita2(request: Request, visit_id: int, db: Session = Depends(get_db)):
    visit = db.query(models.TechnicalVisit2).filter(models.TechnicalVisit2.id == visit_id).first()
    if not visit:
        return RedirectResponse(url="/visitas2", status_code=303)

    clientes = list_customers(db)
    tecnicos = db.query(models.TechnicalAdvisor).all()
    distribuidoras = db.query(models.DistributorPlant).order_by(models.DistributorPlant.name).all()
    temas_calidad = db.query(models.QualityTopic).order_by(models.QualityTopic.name).all()
    calidad_at = db.query(models.QualityAT).order_by(models.QualityAT.name).all()
    # Tareas del plan de acción que generó esta visita: se muestran en el
    # formulario para que el vínculo visita <-> seguimiento sea visible.
    tareas_vinculadas = db.query(models.ActionPlanTracking).filter(
        models.ActionPlanTracking.source_visit_id == visit_id,
        models.ActionPlanTracking.is_active.is_(True),
    ).order_by(models.ActionPlanTracking.id).all()
    return templates.TemplateResponse(request, 'nueva_visita2.html', {
        "request": request,
        "visit": visit,
        "form_data": visit.form_data or {},
        "title": f"Editar Visita V2 #{visit.id}",
        "action_url": f"/visitas2/{visit_id}/editar",
        "active_page": "visitas2",
        "clientes": clientes,
        "tecnicos": tecnicos,
        "distribuidoras": distribuidoras,
        "temas_calidad": temas_calidad,
        "calidad_at": calidad_at,
        "tareas_vinculadas": tareas_vinculadas,
    })


@router.post("/visitas2/{visit_id}/editar")
async def save_editar_visita2(visit_id: int, request: Request, db: Session = Depends(get_db)):
    visit = db.query(models.TechnicalVisit2).filter(models.TechnicalVisit2.id == visit_id).first()
    if not visit:
        return RedirectResponse(url="/visitas2", status_code=303)

    form = await request.form()
    data = safe_json_loads(form.get("form_data"))

    visit.fecha = data.get("fecha", "") or ""
    visit.status = data.get("status", "En proceso")
    visit.form_data = data
    visit.updated_at = datetime.datetime.utcnow()
    db.commit()

    sync_action_plan(db, visit.id, data)
    db.commit()

    return RedirectResponse(url="/visitas2", status_code=303)


@router.post("/cancelar-visita2/{visit_id}")
async def cancel_visita2(visit_id: int, request: Request, db: Session = Depends(get_db)):
    visit = db.query(models.TechnicalVisit2).filter(models.TechnicalVisit2.id == visit_id).first()
    if visit:
        visit.status = "Cancelado"
        fd = visit.form_data or {}
        fd["status"] = "Cancelado"
        visit.form_data = fd
        db.commit()
    return RedirectResponse(url="/visitas2", status_code=303)


@router.post("/eliminar-visita2/{visit_id}")
async def delete_visita2(visit_id: int, request: Request, db: Session = Depends(get_db)):
    visit = db.query(models.TechnicalVisit2).filter(models.TechnicalVisit2.id == visit_id).first()
    if visit:
        from services.trash_service import move_to_trash
        user = getattr(request.state, 'user', None)
        uid = int(user["sub"]) if user else None
        uname = user.get("username") if user else None
        fd = visit.form_data or {}

        # Las tareas del plan de acción que nacieron de esta visita se desactivan
        # junto con ella. Antes quedaban activas y huérfanas: seguían apareciendo
        # en /seguimiento-plan apuntando a una visita que ya no existía, aunque el
        # aviso de borrado prometía eliminarlas. No se borran para poder
        # recuperarlas si se restaura la visita desde la papelera.
        tareas = db.query(models.ActionPlanTracking).filter(
            models.ActionPlanTracking.source_visit_id == visit_id,
            models.ActionPlanTracking.is_active.is_(True),
        ).all()
        for tarea in tareas:
            tarea.is_active = False
            tarea.inactivation_reason = f"Visita #{visit_id} eliminada"
            tarea.inactivated_at = datetime.datetime.now()

        move_to_trash(db, "visit2", visit.id, {
            "id": visit.id,
            "fecha": visit.fecha,
            "status": visit.status,
            "form_data": fd,
            "cliente_nombre": fd.get("cliente_nombre", ""),
            "tareas_ids": [t.id for t in tareas],
        }, uid, uname, f"Visita #{visit.id}")
        db.delete(visit)
        db.commit()
    return RedirectResponse(url="/visitas2?ok=eliminada", status_code=303)
