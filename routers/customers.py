import io
import datetime
from fastapi import APIRouter, Request, Depends, Form, HTTPException, Query
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse, StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from typing import Optional, List
import models
from database import get_db
from services import customer_service, pagination as pg, soft_delete as sd
from services import auditoria, lote
from templating import templates

router = APIRouter(tags=["customers"])


@router.get("/organigrama", response_class=HTMLResponse)
async def organigrama_page(request: Request, db: Session = Depends(get_db)):
    customers = db.query(models.Customer).filter(models.Customer.is_active == True).order_by(models.Customer.name).all()
    return templates.TemplateResponse(request, 'organigrama.html', {
        "request": request,
        "customers": customers,
        "active_page": "organigrama"
    })


@router.get("/clientes/exportar-excel")
async def export_customers_excel(include_inactive: bool = False,
                                 db: Session = Depends(get_db)):
    """Exporta los clientes a Excel.

    Por defecto sólo los activos, para que coincida con el listado, que oculta
    los dados de baja. Antes exportaba todos y el archivo mostraba clientes que
    la pantalla no listaba.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    consulta = db.query(models.Customer)
    if not include_inactive:
        consulta = consulta.filter(models.Customer.is_active.is_(True))
    customers = consulta.order_by(models.Customer.name).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"

    header_font = Font(name="Inter", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1A57DB", end_color="1A57DB", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    headers = ["Nombre", "Planta", "Sector Industrial", "Comercial Asignado",
               "Asesor Técnico", "Diseñador", "Dirección", "Ciudad",
               "Status Comercial", "Status Relevamiento", "Visitas Std",
               "Visitas Premium", "Observaciones"]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row, c in enumerate(customers, 2):
        ws.cell(row=row, column=1, value=c.name)
        ws.cell(row=row, column=2, value=c.plant)
        ws.cell(row=row, column=3, value=c.industry_sector)
        ws.cell(row=row, column=4, value=c.assigned_commercial)
        ws.cell(row=row, column=5, value=c.technical_advisor)
        ws.cell(row=row, column=6, value=c.designer)
        ws.cell(row=row, column=7, value=c.address)
        ws.cell(row=row, column=8, value=c.city)
        ws.cell(row=row, column=9, value=c.status)
        ws.cell(row=row, column=10, value=c.relevamiento_status)
        ws.cell(row=row, column=11, value=c.standard_visits or 0)
        ws.cell(row=row, column=12, value=c.premium_visits or 0)
        ws.cell(row=row, column=13, value=c.observations)

    for col in ws.column_dimensions:
        ws.column_dimensions[col].width = 22

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["G"].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=clientes.xlsx"}
    )


@router.get("/clientes", response_class=HTMLResponse)
async def list_customers(
    request: Request,
    q: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    per_page: int = Query(10, ge=1, le=100),
    db: Session = Depends(get_db)
):
    query = db.query(models.Customer).filter(models.Customer.is_active.is_(True))
    if q:
        query = query.filter(
            (models.Customer.name.ilike(f"%{q}%")) |
            (models.Customer.plant.ilike(f"%{q}%")) |
            (models.Customer.industry_sector.ilike(f"%{q}%")) |
            (models.Customer.city.ilike(f"%{q}%"))
        )
    total = query.count()
    page, per_page, offset = pg.paginate(page, per_page)
    # Stats describe the whole filtered set, not just the visible page. Se piden
    # ANTES de aplicar el recorte: si se aplican limit/offset primero quedan en la
    # consulta y el agregado describe sólo la página visible (y cambia al paginar).
    stats = customer_service.get_customer_stats(db, base_query=query)
    customers = query.order_by(models.Customer.name).offset(offset).limit(per_page).all()
    client_names = db.query(models.ClientName).order_by(models.ClientName.name).all()
    plant_names = db.query(models.PlantName).order_by(models.PlantName.name).all()

    return templates.TemplateResponse(request, 'clientes.html', {
        "request": request, "customers": customers, "stats": stats, "q": q, "active_page": "clientes",
        "advisors": db.query(models.TechnicalAdvisor).order_by(models.TechnicalAdvisor.name).all(),
        "designers": db.query(models.Designer).order_by(models.Designer.name).all(),
        "commercials": db.query(models.CommercialAdvisor).order_by(models.CommercialAdvisor.name).all(),
        "sectors": db.query(models.IndustrialSector).order_by(models.IndustrialSector.name).all(),
        "client_names": client_names,
        "plant_names": plant_names,
        "member_sectors": db.query(models.MemberSector).order_by(models.MemberSector.name).all(),
        "pagination": pg.build_pagination(page, per_page, total),
    })


@router.get("/nuevo-cliente", response_class=HTMLResponse)
async def new_customer_form(request: Request, db: Session = Depends(get_db)):
    advisors = db.query(models.TechnicalAdvisor).order_by(models.TechnicalAdvisor.name).all()
    designers = db.query(models.Designer).order_by(models.Designer.name).all()
    commercials = db.query(models.CommercialAdvisor).order_by(models.CommercialAdvisor.name).all()
    sectors = db.query(models.IndustrialSector).order_by(models.IndustrialSector.name).all()
    client_names = db.query(models.ClientName).order_by(models.ClientName.name).all()
    plant_names = db.query(models.PlantName).order_by(models.PlantName.name).all()
    return templates.TemplateResponse(request, 'nuevo_cliente.html', {
        "request": request,
        "advisors": advisors,
        "designers": designers,
        "commercials": commercials,
        "sectors": sectors,
        "client_names": client_names,
        "plant_names": plant_names,
        "member_sectors": db.query(models.MemberSector).order_by(models.MemberSector.name).all(),
        "active_page": "clientes"
    })


@router.get("/api/cliente-plantas/{customer_name}")
async def get_customer_plants(customer_name: str, db: Session = Depends(get_db)):
    return customer_service.get_customer_plants(db, customer_name)


@router.get("/api/plantas-lineas/{customer_name}/{plant_name}")
async def get_plant_lines(customer_name: str, plant_name: str, db: Session = Depends(get_db)):
    return customer_service.get_plant_lines(db, customer_name, plant_name)


@router.get("/api/linea-maquinas/{customer_name}/{plant_name}/{line_name}")
async def get_line_machines(customer_name: str, plant_name: str, line_name: str, db: Session = Depends(get_db)):
    return customer_service.get_line_machines(db, customer_name, plant_name, line_name)


@router.get("/api/planta-maquinas/{customer_name}/{plant_name}")
async def get_plant_machines(customer_name: str, plant_name: str, db: Session = Depends(get_db)):
    return customer_service.get_plant_machines(db, customer_name, plant_name)


@router.get("/api/cliente-contactos/{customer_id}")
async def get_cliente_contactos(customer_id: int, db: Session = Depends(get_db)):
    customer = customer_service.get_customer(db, customer_id)
    names = list({m["name"] for m in (customer.client_team or []) if m.get("name")})
    return sorted(names)


@router.get("/api/cliente-team/{customer_id}")
async def get_cliente_team(customer_id: int, db: Session = Depends(get_db)):
    customer = customer_service.get_customer(db, customer_id)
    return customer.client_team or []


@router.post("/api/contact-name/quick-create")
async def quick_create_contact_name(name: str = Form(...), db: Session = Depends(get_db)):
    existing = db.query(models.ContactName).filter(func.upper(models.ContactName.name) == name.strip().upper()).first()
    if existing:
        return {"id": existing.id, "name": existing.name}
    contact = models.ContactName(name=name.strip().upper())
    db.add(contact)
    db.commit()
    db.refresh(contact)
    return {"id": contact.id, "name": contact.name}


@router.post("/nuevo-cliente")
async def create_customer(
    request: Request,
    name: str = Form(...), plant: str = Form(...), industry_sector: str = Form(...),
    assigned_commercial: str = Form(...), technical_advisor: str = Form(...),
    designer: str = Form(...), address: str = Form(...), city: str = Form(...),
    status: str = Form("Inactivo"), relevamiento_status: Optional[str] = Form(None),
    member_name: List[str] = Form([]), member_role: List[str] = Form([]),
    member_phone: List[str] = Form([]),
    member_reports_to: List[str] = Form([]),
    member_sector: List[str] = Form([]),
    observations: str = Form(""),
    db: Session = Depends(get_db)
):
    creado = customer_service.create_customer(db, name, plant, industry_sector, assigned_commercial,
                                      technical_advisor, designer, address, city,
                                      status, relevamiento_status,
                                      member_name, member_role, member_phone, member_reports_to, member_sector, observations)
    # El historial se registra después del alta: si el usuario no tiene permiso o
    # los datos no validan, no queda un registro de algo que no pasó.
    if creado is not None:
        auditoria.registrar(db, "customer", getattr(creado, "id", 0), "crear",
                            user=getattr(request.state, "user", None),
                            entity_label=name)
    return RedirectResponse(url="/clientes", status_code=303)


@router.get("/editar-cliente/{customer_id}", response_class=HTMLResponse)
async def edit_customer_form(customer_id: int, request: Request, db: Session = Depends(get_db)):
    # If not an AJAX request, redirect to the clientes page (the edit modal opens from there)
    if request.headers.get("X-Requested-With") != "XMLHttpRequest":
        return RedirectResponse(url="/clientes", status_code=302)
    customer = customer_service.get_customer(db, customer_id)
    advisors = db.query(models.TechnicalAdvisor).order_by(models.TechnicalAdvisor.name).all()
    designers = db.query(models.Designer).order_by(models.Designer.name).all()
    commercials = db.query(models.CommercialAdvisor).order_by(models.CommercialAdvisor.name).all()
    sectors = db.query(models.IndustrialSector).order_by(models.IndustrialSector.name).all()
    client_names = db.query(models.ClientName).order_by(models.ClientName.name).all()
    plant_names = db.query(models.PlantName).order_by(models.PlantName.name).all()
    return templates.TemplateResponse(request, '_editar_cliente_partial.html', {
        "request": request,
        "customer": customer,
        "advisors": advisors,
        "designers": designers,
        "commercials": commercials,
        "sectors": sectors,
        "client_names": client_names,
        "plant_names": plant_names,
        "member_sectors": db.query(models.MemberSector).order_by(models.MemberSector.name).all(),
    })


@router.post("/editar-cliente/{customer_id}")
async def update_customer(
    request: Request,
    customer_id: int,
    name: str = Form(...), plant: str = Form(...), industry_sector: str = Form(...),
    assigned_commercial: str = Form(...), technical_advisor: str = Form(...),
    designer: str = Form(...), address: str = Form(...), city: str = Form(...),
    status: str = Form(...), relevamiento_status: Optional[str] = Form(None),
    member_name: List[str] = Form([]), member_role: List[str] = Form([]),
    member_phone: List[str] = Form([]),
    member_reports_to: List[str] = Form([]),
    member_sector: List[str] = Form([]),
    observations: str = Form(""),
    db: Session = Depends(get_db)
):
    customer = customer_service.get_customer(db, customer_id)
    if customer and not customer.is_active:
        raise HTTPException(status_code=403, detail="No se puede editar un cliente inactivo")
    # Foto del estado previo: sin esto no se puede decir qué cambió.
    antes = auditoria.snapshot(customer)
    customer_service.update_customer(db, customer_id, name, plant, industry_sector,
                                       assigned_commercial, technical_advisor, designer,
                                       address, city, status, relevamiento_status,
                                       member_name, member_role, member_phone, member_reports_to, member_sector, observations)
    # Se relee el registro para comparar contra la foto previa y registrar sólo
    # los campos que efectivamente cambiaron.
    despues = auditoria.snapshot(customer_service.get_customer(db, customer_id))
    cambios = auditoria.diferencias(antes, despues)
    if cambios:
        auditoria.registrar(db, "customer", customer_id, "editar",
                            user=getattr(request.state, "user", None),
                            entity_label=name, changes=cambios)
    return RedirectResponse(url="/clientes", status_code=303)


@router.get("/cliente-reporte/{customer_id}", response_class=HTMLResponse)
async def customer_report(customer_id: int, request: Request, db: Session = Depends(get_db)):
    customer = customer_service.get_customer(db, customer_id)
    machines = db.query(models.Machine).filter(models.Machine.customer_id == customer_id).order_by(models.Machine.brand, models.Machine.model).all()
    return templates.TemplateResponse(request, 'cliente_reporte.html', {
        "request": request,
        "customer": customer,
        "machines": machines,
        "visits": [],
        "generation_date": datetime.datetime.now(),
        "active_page": "clientes"
    })


# ── Acciones en lote ────────────────────────────────────────────────────────
# Desactivar veinte clientes de a uno, con su motivo cada vez, es la clase de tarea
# que hace que la gente deje de usar la herramienta. La lógica compartida —y el
# tope de identificadores— están en `services/lote.py`.


@router.post("/clientes/accion-en-lote")
async def clientes_accion_en_lote(request: Request, db: Session = Depends(get_db)):
    """Aplica una acción a varios clientes a la vez.

    Devuelve un resumen con cuántos se aplicaron y cuáles no, en lugar de un simple
    "ok": si tres de veinte fallaron, hay que poder decirlo.
    """
    datos = await request.json()
    ids = datos.get("ids") or []
    accion = (datos.get("accion") or "").strip()
    motivo = (datos.get("motivo") or "").strip()

    def error(mensaje, codigo=400):
        return JSONResponse(status_code=codigo, content={"detail": mensaje})

    if accion not in ("desactivar", "reactivar"):
        return error("Acción no válida")

    user = getattr(request.state, "user", None)
    ids_limpios, falla = lote.validar_peticion(ids, accion, motivo, user)
    if falla:
        return error(falla["mensaje"], falla["codigo"])

    aplicados, omitidos = lote.aplicar(
        db, models.Customer, ids_limpios, accion, motivo, user,
        etiqueta=lambda c: c.name, registrar="customer")

    return JSONResponse({
        "aplicados": aplicados,
        "omitidos": omitidos,
        "accion": accion,
    })


@router.post("/desactivar-cliente/{customer_id}")
async def deactivate_customer(customer_id: int, request: Request, db: Session = Depends(get_db)):
    data = await request.json()
    reason = data.get("reason", "")
    user = getattr(request.state, 'user', None)
    uid = int(user["sub"]) if user else None
    resultado = sd.soft_deactivate(db, models.Customer, customer_id, reason, uid)
    cliente = db.query(models.Customer).filter(models.Customer.id == customer_id).first()
    auditoria.registrar(db, "customer", customer_id, "desactivar",
                        user=user, entity_label=cliente.name if cliente else None,
                        changes={"motivo": {"antes": "", "despues": reason or "sin motivo"}})
    return resultado


@router.post("/reactivar-cliente/{customer_id}")
async def reactivate_customer(customer_id: int, request: Request, db: Session = Depends(get_db)):
    user = getattr(request.state, 'user', None)
    role = (user.get("role") or "") if user else ""
    if role != "admin":
        raise HTTPException(status_code=403, detail="Solo el admin puede reactivar")
    resultado = sd.soft_reactivate(db, models.Customer, customer_id)
    cliente = db.query(models.Customer).filter(models.Customer.id == customer_id).first()
    auditoria.registrar(db, "customer", customer_id, "reactivar",
                        user=user, entity_label=cliente.name if cliente else None)
    return resultado
